import io
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import HttpResponse, Http404
from .models import Resume, ResumeTemplate, Section, SectionEntry
from .forms import ResumeForm, SectionForm, SectionEntryForm, SectionEntryFormSet


# ── Home ──────────────────────────────────────────────────────────────────────

def home(request):
    templates = ResumeTemplate.objects.filter(is_active=True)[:6]
    from announcements.models import Announcement
    news = Announcement.objects.filter(is_published=True).order_by('-created_at')[:3]
    return render(request, 'home.html', {'templates': templates, 'news': news})


# ── Template catalog ──────────────────────────────────────────────────────────

def template_catalog(request):
    templates = ResumeTemplate.objects.filter(is_active=True)
    return render(request, 'resumes/template_catalog.html', {'templates': templates})


# ── Resume list ───────────────────────────────────────────────────────────────

@login_required
def resume_list(request):
    resumes = Resume.objects.filter(owner=request.user).select_related('template')
    return render(request, 'resumes/resume_list.html', {'resumes': resumes})


# ── Resume create ─────────────────────────────────────────────────────────────

@login_required
def resume_create(request):
    template_id = request.GET.get('template')
    initial = {}
    if template_id:
        try:
            tmpl = ResumeTemplate.objects.get(pk=template_id, is_active=True)
            initial['template'] = tmpl
        except ResumeTemplate.DoesNotExist:
            pass

    if request.method == 'POST':
        form = ResumeForm(request.POST, request.FILES)
        if form.is_valid():
            resume = form.save(commit=False)
            resume.owner = request.user
            resume.save()
            messages.success(request, 'Резюме створено! Тепер додайте секції.')
            return redirect('resume_edit', pk=resume.pk)
    else:
        form = ResumeForm(initial=initial)
    return render(request, 'resumes/resume_form.html', {'form': form, 'action': 'create'})


# ── Resume edit ───────────────────────────────────────────────────────────────

@login_required
def resume_edit(request, pk):
    resume = get_object_or_404(Resume, pk=pk, owner=request.user)
    if request.method == 'POST':
        form = ResumeForm(request.POST, request.FILES, instance=resume)
        if form.is_valid():
            form.save()
            messages.success(request, 'Зміни збережено.')
            return redirect('resume_edit', pk=resume.pk)
    else:
        form = ResumeForm(instance=resume)
    sections = resume.sections.prefetch_related('entries').all()
    return render(request, 'resumes/resume_edit.html', {
        'form': form,
        'resume': resume,
        'sections': sections,
    })


# ── Resume detail / preview ───────────────────────────────────────────────────

@login_required
def resume_detail(request, pk):
    resume = get_object_or_404(Resume, pk=pk, owner=request.user)
    sections = resume.sections.prefetch_related('entries').all()
    return render(request, 'resumes/resume_detail.html', {
        'resume': resume,
        'sections': sections,
    })


# ── Resume delete ─────────────────────────────────────────────────────────────

@login_required
def resume_delete(request, pk):
    resume = get_object_or_404(Resume, pk=pk, owner=request.user)
    if request.method == 'POST':
        resume.delete()
        messages.success(request, 'Резюме видалено.')
        return redirect('resume_list')
    return render(request, 'resumes/resume_confirm_delete.html', {'resume': resume})


# ── Resume clone ──────────────────────────────────────────────────────────────

@login_required
def resume_clone(request, pk):
    resume = get_object_or_404(Resume, pk=pk, owner=request.user)
    new_resume = resume.clone()
    messages.success(request, f'Резюме "{resume.title}" скопійовано.')
    return redirect('resume_edit', pk=new_resume.pk)


# ── Section add ───────────────────────────────────────────────────────────────

@login_required
def section_add(request, resume_pk):
    resume = get_object_or_404(Resume, pk=resume_pk, owner=request.user)
    if request.method == 'POST':
        form = SectionForm(request.POST)
        if form.is_valid():
            section = form.save(commit=False)
            section.resume = resume
            section.save()
            messages.success(request, 'Секцію додано.')
            return redirect('resume_edit', pk=resume.pk)
    else:
        form = SectionForm()
    return render(request, 'resumes/section_form.html', {
        'form': form, 'resume': resume, 'action': 'add'
    })


# ── Section edit (with entries formset) ──────────────────────────────────────

@login_required
def section_edit(request, resume_pk, section_pk):
    resume = get_object_or_404(Resume, pk=resume_pk, owner=request.user)
    section = get_object_or_404(Section, pk=section_pk, resume=resume)
    if request.method == 'POST':
        form = SectionForm(request.POST, instance=section)
        formset = SectionEntryFormSet(request.POST, instance=section)
        if form.is_valid() and formset.is_valid():
            form.save()
            formset.save()
            messages.success(request, 'Секцію оновлено.')
            return redirect('resume_edit', pk=resume.pk)
    else:
        form = SectionForm(instance=section)
        formset = SectionEntryFormSet(instance=section)
    return render(request, 'resumes/section_form.html', {
        'form': form,
        'formset': formset,
        'resume': resume,
        'section': section,
        'action': 'edit',
    })


# ── Section delete ────────────────────────────────────────────────────────────

@login_required
def section_delete(request, resume_pk, section_pk):
    resume = get_object_or_404(Resume, pk=resume_pk, owner=request.user)
    section = get_object_or_404(Section, pk=section_pk, resume=resume)
    if request.method == 'POST':
        section.delete()
        messages.success(request, 'Секцію видалено.')
        return redirect('resume_edit', pk=resume.pk)
    return render(request, 'resumes/section_confirm_delete.html', {
        'section': section, 'resume': resume
    })


# ── Export PDF ────────────────────────────────────────────────────────────────

@login_required
def resume_export_pdf(request, pk):
    resume = get_object_or_404(Resume, pk=pk, owner=request.user)
    sections = resume.sections.prefetch_related('entries').all()
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
        from reportlab.lib.enums import TA_CENTER, TA_LEFT

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer, pagesize=A4,
            rightMargin=2*cm, leftMargin=2*cm,
            topMargin=2*cm, bottomMargin=2*cm
        )
        styles = getSampleStyleSheet()
        story = []

        title_style = ParagraphStyle('Title', parent=styles['Title'],
                                      fontSize=22, spaceAfter=4, textColor=colors.HexColor('#2563eb'))
        subtitle_style = ParagraphStyle('Subtitle', parent=styles['Normal'],
                                         fontSize=13, spaceAfter=2, textColor=colors.HexColor('#475569'))
        heading_style = ParagraphStyle('Heading', parent=styles['Heading2'],
                                        fontSize=13, spaceBefore=12, spaceAfter=4,
                                        textColor=colors.HexColor('#1e40af'),
                                        borderPad=(0, 0, 2, 0))
        body_style = ParagraphStyle('Body', parent=styles['Normal'], fontSize=10, leading=14)
        entry_title_style = ParagraphStyle('EntryTitle', parent=styles['Normal'],
                                            fontSize=11, fontName='Helvetica-Bold')
        meta_style = ParagraphStyle('Meta', parent=styles['Normal'],
                                     fontSize=9, textColor=colors.HexColor('#64748b'))

        # Header
        story.append(Paragraph(resume.full_name or resume.title, title_style))
        if resume.desired_position:
            story.append(Paragraph(resume.desired_position, subtitle_style))

        # Contact line
        contacts = ' | '.join(filter(None, [resume.email, resume.phone, resume.address, resume.website]))
        if contacts:
            story.append(Paragraph(contacts, meta_style))
        story.append(HRFlowable(width='100%', thickness=2, color=colors.HexColor('#2563eb'), spaceAfter=8))

        if resume.summary:
            story.append(Paragraph('Про себе', heading_style))
            story.append(Paragraph(resume.summary.replace('\n', '<br/>'), body_style))

        for section in sections:
            story.append(Paragraph(section.title, heading_style))
            story.append(HRFlowable(width='100%', thickness=0.5, color=colors.HexColor('#cbd5e1'), spaceAfter=4))
            for entry in section.entries.all():
                if entry.subtitle:
                    story.append(Paragraph(entry.subtitle, entry_title_style))
                meta_parts = []
                if entry.date_start or entry.date_end:
                    meta_parts.append(f"{entry.date_start}–{entry.date_end}")
                if entry.location:
                    meta_parts.append(entry.location)
                if meta_parts:
                    story.append(Paragraph(' | '.join(meta_parts), meta_style))
                if entry.description:
                    story.append(Paragraph(entry.description.replace('\n', '<br/>'), body_style))
                story.append(Spacer(1, 6))

        doc.build(story)
        buffer.seek(0)
        filename = f"resume_{resume.pk}.pdf"
        response = HttpResponse(buffer, content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response

    except ImportError:
        messages.error(request, 'Бібліотека reportlab не встановлена. Встановіть: pip install reportlab')
        return redirect('resume_detail', pk=pk)


# ── Export DOCX ───────────────────────────────────────────────────────────────

@login_required
def resume_export_docx(request, pk):
    resume = get_object_or_404(Resume, pk=pk, owner=request.user)
    sections = resume.sections.prefetch_related('entries').all()
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()

        # Styles
        style = doc.styles['Normal']
        style.font.name = 'Calibri'
        style.font.size = Pt(11)

        # Title
        title_p = doc.add_heading(resume.full_name or resume.title, level=0)
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in title_p.runs:
            run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

        if resume.desired_position:
            pos_p = doc.add_paragraph(resume.desired_position)
            pos_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pos_p.runs[0].font.size = Pt(13)

        # Contacts
        contacts = ' | '.join(filter(None, [resume.email, resume.phone, resume.address, resume.website]))
        if contacts:
            cp = doc.add_paragraph(contacts)
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.runs[0].font.size = Pt(9)
            cp.runs[0].font.color.rgb = RGBColor(0x64, 0x74, 0x8b)

        doc.add_paragraph('─' * 80)

        if resume.summary:
            doc.add_heading('Про себе', level=2)
            doc.add_paragraph(resume.summary)

        for section in sections:
            doc.add_heading(section.title, level=2)
            for entry in section.entries.all():
                if entry.subtitle:
                    p = doc.add_paragraph()
                    run = p.add_run(entry.subtitle)
                    run.bold = True
                    run.font.size = Pt(11)
                meta_parts = []
                if entry.date_start or entry.date_end:
                    meta_parts.append(f"{entry.date_start}–{entry.date_end}")
                if entry.location:
                    meta_parts.append(entry.location)
                if meta_parts:
                    mp = doc.add_paragraph(' | '.join(meta_parts))
                    mp.runs[0].font.size = Pt(9)
                    mp.runs[0].font.color.rgb = RGBColor(0x64, 0x74, 0x8b)
                if entry.description:
                    doc.add_paragraph(entry.description)

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        filename = f"resume_{resume.pk}.docx"
        response = HttpResponse(
            buffer,
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response

    except ImportError:
        messages.error(request, 'Бібліотека python-docx не встановлена. Встановіть: pip install python-docx')
        return redirect('resume_detail', pk=pk)
